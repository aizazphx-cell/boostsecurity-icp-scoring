"""Build BoostSecurity_ICP_Scoring.xlsx from results.csv + signals.csv (3 tabs: Scores, Signals, Model)."""
import csv, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = pathlib.Path(__file__).parent
wb = Workbook()
BOLD = Font(bold=True)
HEAD = PatternFill("solid", fgColor="DDE3EA")
TIER = {"Tier 1": "C6EFCE", "Tier 2": "FFEB9C", "Tier 3": "F2F2F2", "DQ": "FFC7CE"}


def sheet(ws, rows, widths):
    ws.append(list(rows[0]))
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD
    for r in rows:
        ws.append([float(v) if v.replace(".", "", 1).isdigit() else v for v in r.values()])
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"


ws = wb.active
ws.title = "Scores"
res = list(csv.DictReader(open(ROOT / "results.csv")))
sheet(ws, res, [22, 7, 9, 8, 30, 7, 7, 7, 7, 7, 7, 7, 70])
for row in ws.iter_rows(min_row=2):
    row[3].fill = PatternFill("solid", fgColor=TIER[row[3].value])

ws2 = wb.create_sheet("Signals")
sig = list(csv.DictReader(open(ROOT / "signals.csv")))
sheet(ws2, sig, [22, 18, 7, 9, 7, 70, 60])
for row in ws2.iter_rows(min_row=2):
    if row[6].value:
        row[6].hyperlink = row[6].value
        row[6].font = Font(color="0563C1", underline="single")

ws3 = wb.create_sheet("Model")
model = [
    ("Criterion", "Weight", "Full (1) / partial (0.5)", "Source", "Reliability"),
    ("C1 AppSec/DevSecOps hiring", 20, "AppSec/ProdSec/DevSecOps role / generic security role", "Greenhouse, Ashby, Lever public APIs", "High detect, medium intent"),
    ("C2 Pipeline risk", 15, "Exploitable workflow findings in 3+ / 1-2 sampled repos", "BoostSecurity poutine (open source) over the 30-repo sample", "High detect (static, deterministic); medium meaning: public repos only, patterns not confirmed exploits"),
    ("C3 Supply-chain exposure", 15, "High/critical GHSA or supply-chain incident <12m / 12-24m", "GitHub Security Advisories API + incident post-mortems", "High, but sparse"),
    ("C4 New security leader", 15, "Started <6m / 6-12m", "Press releases, news, leadership pages", "Low-medium: lagging, often unannounced"),
    ("C5 Firmographic fit", 15, "Software B2B 200-5,000 / 100-199 or 5,001-15,000", "Company site, 10-K, headcount estimates", "Medium"),
    ("C6 Compliance pressure", 10, "Trust center + FedRAMP or regulated customers / SOC 2 only", "trust.* pages, FedRAMP Marketplace", "High presence, low timing"),
    ("C7 Recent funding", 10, "Series B+ <12m / 12-24m or IPO <24m", "Company press releases", "Medium-high"),
    ("Gates (score -> 0)", "", "Competitor; no commercial software product; <100 employees; non-software enterprise", "Company website", ""),
    ("Score", "", "Gate x sum(weight x value). Tier 1 >= 65 and one full-strength timing trigger (C1, C3 or C4 = 1), Tier 2 40-64, Tier 3 < 40", "", ""),
]
for r in model:
    ws3.append(r)
for c in ws3[1]:
    c.font, c.fill = BOLD, HEAD
for i, w in enumerate([28, 8, 55, 45, 35]):
    ws3.column_dimensions[chr(65 + i)].width = w

wb.save(ROOT / "BoostSecurity_ICP_Scoring.xlsx")
print("saved")
